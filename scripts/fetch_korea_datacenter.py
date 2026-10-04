"""Poll public BLCM exports. Authentication-free; original workbooks stay local.

Snapshots are request-driven exports, not a guaranteed monthly feed. Only the
nationwide permit/start families are collected; title and data dates are checked.
"""
import argparse
import copy
import json
import re
import sys
from io import BytesIO
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from common import collection_date, UA
from korea_datacenter import BASELINE, OUTPUT, ROOT, PUBLIC_URL, build_documents, merge_observations, parse_public_rows, write_documents

CACHE = ROOT / '.cache/korea_datacenter'
STATE = OUTPUT / 'observations.json'


def listing(html):
    found = []
    for row in BeautifulSoup(html, 'html.parser').select('tbody tr'):
        cells, links = row.select('td'), row.select('a')
        if len(cells) < 3 or not links:
            continue
        title = links[0].get_text(strip=True)
        kind = 'permits' if '전국 건축물 허가 현황' in title else 'starts' if '건축물 착공 현황' in title else None
        if not kind:
            continue
        file_link = next((a.get('href', '') for a in links if 'fn_egov_downFile' in a.get('href', '')), '')
        match = re.search(r"fn_egov_downFile\('([^']+)',\s*'([^']+)'", file_link)
        if not match:
            continue
        found.append({'id': cells[0].get_text(strip=True), 'title': title, 'published': cells[2].get_text(strip=True),
                      'kind': kind, 'url': PUBLIC_URL, 'file_id': match[1], 'file_sn': match[2]})
    return found


def declared_period(title):
    if '전국 건축물 허가 현황' not in title:
        return None
    m = re.search(r'(\d{4})년\s*(\d{1,2})월\s*[~～]\s*(?:(\d{4})년\s*)?(\d{1,2})월', title)
    if not m:
        return None
    import calendar
    a, b, c, d = map(int, [m[1], m[2], m[3] or m[1], m[4]])
    return f'{a:04}-{b:02}-01', f'{c:04}-{d:02}-{calendar.monthrange(c, d)[1]:02}'


def read_export(content, source):
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if '상세리스트' not in wb.sheetnames:
            raise ValueError('Expected 상세리스트 sheet missing')
        rows = wb['상세리스트'].iter_rows(values_only=True)
        header = next(rows)
        # Filter incrementally rather than retaining a nationwide multi-million-cell grid.
        matching, count = [], 0
        for row in rows:
            count += 1
            if re.search(r'데이터\s*센터|data\s*center|\bA?IDC\b', ' '.join(str(x) for x in row if x is not None), re.I):
                matching.append(row)
        observations, stats = parse_public_rows(header, matching, source, starts=source['kind'] == 'starts')
        stats['rows'] = count
        if not count or not stats['date_min']:
            raise ValueError('Export contains no dated data-centre observations')
        period = declared_period(source['title'])
        if period and not (period[0] <= stats['date_min'] <= stats['date_max'] <= period[1]):
            raise ValueError('Permit dates disagree with export title; refusing snapshot')
        stats['declared_period'] = list(period) if period else None
        return observations, stats
    finally:
        wb.close()


def collect(max_pages=6):
    baseline = json.loads(BASELINE.read_text(encoding='utf8'))
    state = json.loads(STATE.read_text(encoding='utf8')) if STATE.exists() else {'observations': [], 'manifest': []}
    session = requests.Session()
    session.headers.update(UA)
    found = []
    for page in range(1, max_pages + 1):
        r = session.get(PUBLIC_URL, params={'pageIndex': page}, timeout=(10, 45))
        r.raise_for_status()
        found.extend(listing(r.content.decode('utf8')))
        if sum(s['kind'] == 'permits' for s in found) >= 2 and any(s['kind'] == 'starts' for s in found):
            break
    selected = []
    for kind, limit in [('permits', 2), ('starts', 1)]:
        family = sorted((s for s in found if s['kind'] == kind), key=lambda s: int(s['id']), reverse=True)
        if not family:
            raise ValueError(f'No public {kind} exports found')
        selected.extend(family[:limit])
    CACHE.mkdir(parents=True, exist_ok=True)
    manifest = {s['id']: s for s in state['manifest']}
    observations = copy.deepcopy(state['observations'])
    errors = []
    for source in selected:
        if source['id'] in manifest:
            continue
        try:
            path = CACHE / (source['id'] + '.xlsx')
            if not path.exists():
                r = session.get('https://blcm.go.kr/cmm/fms/FileDown.do',
                    params={'atchFileId': source['file_id'], 'fileSn': source['file_sn']}, timeout=(15, 120))
                r.raise_for_status()
                if not r.content.startswith(b'PK'):
                    raise ValueError('Download did not return an XLSX workbook')
                path.write_bytes(r.content)
            public_source = {k: v for k, v in source.items() if k not in ['file_id','file_sn']}
            incoming, stats = read_export(path.read_bytes(), public_source)
            observations = merge_observations(observations, incoming)
            manifest[source['id']] = {**public_source, **stats}
            print(f"Imported export {source['id']}: {stats['centres']} permit/phase observations")
        except Exception as error:
            errors.append(f"{source['id']}: {type(error).__name__} · {str(error)[:140]}")
    status = {'checked': collection_date(), 'ok': not errors,
              'message': '공개 목록 확인 완료 · 동일 기준 자료의 월간 제공은 보장되지 않음' if not errors else '일부 자료 수집 실패 · 마지막 성공 자료 유지: ' + '; '.join(errors)}
    OUTPUT.mkdir(parents=True, exist_ok=True)
    state = {'observations': observations, 'manifest': sorted(manifest.values(), key=lambda s: int(s['id'])), 'last_check': status}
    STATE.write_text(json.dumps(state, ensure_ascii=False, indent=1) + '\n', encoding='utf8')
    write_documents(build_documents(baseline, observations, state['manifest'], status))
    if errors:
        raise RuntimeError('; '.join(errors))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--offline', action='store_true', help='Rebuild from saved observations without claiming a fresh collection')
    args = parser.parse_args()
    if args.offline:
        baseline = json.loads(BASELINE.read_text(encoding='utf8'))
        state = json.loads(STATE.read_text(encoding='utf8')) if STATE.exists() else {'observations': [], 'manifest': []}
        write_documents(build_documents(baseline, state['observations'], state['manifest'], state.get('last_check')))
    else:
        try:
            collect()
        except Exception as error:
            # Listing failures also remain visible, with prior data/fetched dates unchanged.
            from common import record_fetch_failure
            for name in ['dc_permits','dc_starts','dc_started_area','dc_started_capacity','dc_facilities']:
                record_fetch_failure('datacenter', name, error)
            print(f'Collection failed: {type(error).__name__}: {str(error)[:240]}', file=sys.stderr)
            return 1
    return 0


if __name__ == '__main__':
    raise SystemExit(main())

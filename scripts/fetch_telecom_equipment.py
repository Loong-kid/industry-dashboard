"""Public telecom equipment accounts, separate from operator investments."""
import argparse
from concurrent.futures import ThreadPoolExecutor
import datetime as dt
import hashlib
import json
from pathlib import Path
import time

import requests

from fetch_cpu_companies import save, reference
from fetch_game_fx import SOURCE, parse_rates, build_snapshot
from fetch_telecom import collect_dart, document, load, merge_rows
from telecom_sources import sec_accounts, dart_accounts, reporting_period

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'data/telecom'
CACHE=ROOT/'.cache/telecom/equipment'
META=ROOT/'manual/telecom/equipment.json'
METRICS=['revenue','profit','capex']


def fetch(session,url,offline=False):
    path=CACHE/(hashlib.sha256(url.encode()).hexdigest()[:18]+'.bin')
    if offline:return path.read_bytes()
    r=session.get(url,timeout=40);r.raise_for_status()
    path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(r.content)
    return r.content


def ifrs_annual(payload,meta,asof):
    if int(payload['cik'])!=meta['cik']:raise ValueError('issuer mismatch')
    tags={'revenue':['RevenueFromContractsWithCustomers','Revenue'],
          'profit':['ProfitLossFromOperatingActivities'],
          'capex':['PurchaseOfPropertyPlantAndEquipmentClassifiedAsInvestingActivities'] if meta['currency']=='SEK' else ['PurchaseOfPropertyPlantAndEquipmentIntangibleAssetsOtherThanGoodwillInvestmentPropertyAndOtherNoncurrentAssets']}
    rows={}
    for metric,names in tags.items():
        selected={}
        for priority,tag in enumerate(names):
            for f in payload['facts']['ifrs-full'].get(tag,{}).get('units',{}).get(meta['currency'],[]):
                if not f.get('start') or not '2020-01-01'<=f['end']<=asof or f['filed']>asof:continue
                if not 350<=(dt.date.fromisoformat(f['end'])-dt.date.fromisoformat(f['start'])).days+1<=380:continue
                rank=(f['filed'],-priority,f['accn'])
                if f['end'] not in selected or rank>selected[f['end']]['rank']:selected[f['end']]=dict(f,tag='ifrs-full:'+tag,rank=rank)
        for end,f in selected.items():
            row=rows.setdefault(end,dict(reporting_period(int(end[:4])),start=f['start'],date=end))
            row[metric]=f['val']/1e6;row['refs'][metric]=dict(reference(f,meta['cik']),calculation='reported annual')
    return list(rows.values())


def collect(offline=False,full=False):
    asof=dt.datetime.now(dt.timezone(dt.timedelta(hours=9))).date().isoformat()
    meta=load(META,{});state=load(OUT/'equipment_state.json',dict(companies={}))
    session=requests.Session();session.headers['User-Agent']='industry-dashboard personal research wkorotk@gmail.com'
    def korean_job(key):
        return collect_dart(meta[key],state['companies'].get(key,{}),requests.Session(),asof,full)
    with ThreadPoolExecutor(max_workers=3) as pool:
        jobs={key:pool.submit(korean_job,key) for key,m in meta.items() if m.get('corp_code')} if not offline else {}
        for key,m in meta.items():
            prior=state['companies'].get(key,{})
            try:
                if m.get('corp_code'):
                    fresh=dict(prior,rows=dart_accounts(prior['reports'],asof)) if offline else jobs[key].result()
                elif m.get('cik'):
                    payload=json.loads(fetch(session,f'https://data.sec.gov/api/xbrl/companyfacts/CIK{m["cik"]:010d}.json',offline))
                    fresh=dict(rows=ifrs_annual(payload,m,asof) if m.get('ifrs') else sec_accounts(payload,m,asof))
                    time.sleep(.2)
                else:fresh=dict(rows=prior.get('rows',[]))
                state['companies'][key]=dict(prior,**fresh,checked=asof,fetched=asof,status='ok')
                print(key,len(fresh['rows']),'periods',flush=True)
            except (requests.RequestException,ValueError,KeyError,OSError) as error:
                state['companies'][key]=dict(prior,checked=asof,status='최신 조회 미완료 · 기존 자료 유지 ('+type(error).__name__+')')
                print(key,'retained previous accounts:',type(error).__name__,flush=True)
    save(OUT/'equipment_state.json',state)
    # Official public workbook is retained in the artifact for format review.
    if not offline:
        import openpyxl,io
        url='https://www.nokia.com/system/files/2026-10/nokia_tables_2026_q3.xlsx'
        try:
            raw=fetch(session,url);w=openpyxl.load_workbook(io.BytesIO(raw),data_only=True)
            for s in w:
                print('Nokia workbook:',s.title,s.max_row,s.max_column,flush=True)
                for row in list(s)[:12]:print([c.value for c in row[:12]],flush=True)
        except (requests.RequestException,ValueError):print('Nokia direct workbook unavailable; official PDF fallback required',flush=True)
    return state


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--offline',action='store_true');p.add_argument('--full-history',action='store_true');a=p.parse_args();collect(a.offline,a.full_history)

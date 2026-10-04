"""Source semantics, revisions, missing observations and collection failure isolation."""
import io
import json
from datetime import datetime
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

import common
import fetch_copper_construction as fc


def workbook_bytes(workbook):
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class ConstructionTests(unittest.TestCase):
    def test_nbs_empty_is_not_zero_and_future_is_rejected(self):
        payload = {"success": True, "data": [
            {"code": "202608MM", "values": [{"_id": "known", "value": "-24.8"}, {"_id": "other", "value": "99"}]},
            {"code": "202607MM", "values": [{"_id": "known", "value": ""}]},
            {"code": "202606MM", "values": [{"_id": "known", "value": "0"}]},
            {"code": "202612MM", "values": [{"_id": "known", "value": "999"}]},
        ]}
        self.assertEqual(fc.nbs_series(payload, ["known"], "2026-10-05"),
                         {"known": [("2026-06-30", 0.0), ("2026-08-31", -24.8)]})
        with self.assertRaises(ValueError):
            fc.nbs_series({"success": False, "data": []}, ["known"])
        with self.assertRaises(ValueError):
            fc.nbs_series(payload, ["missing"])

    def test_revision_merge_preserves_long_history(self):
        self.assertEqual(fc.merged([["2000-02-29", 1], ["2026-08-31", 2]], [["2026-08-31", 3]]),
                         [["2000-02-29", 1], ["2026-08-31", 3]])

    def test_pmi_components_not_manufacturing_or_all_nonmanufacturing(self):
        html = """<title>2026年9月中国采购经理指数运行情况</title><p>2026/09/30</p>
          <p>新订单指数为50.5%；非制造业新订单指数为46.5%；
          建筑业商务活动指数为<span>50.3</span>%，建筑业新订单指数为45.7%。</p>"""
        dt, values, source = fc.parse_pmi(html, "https://www.stats.gov.cn/a.html")
        self.assertEqual(dt, "2026-09-30")
        self.assertEqual(values, {fc.ACTIVITY: 50.3, fc.ORDERS: 45.7})
        self.assertEqual(source["published"], "2026-09-30")
        with self.assertRaises(ValueError):
            fc.parse_pmi(html.replace("建筑业新订单指数", "服务业新订单指数"), source["url"])

    def test_migrated_url_and_january_publication_do_not_relabel_december(self):
        html = """<title>12月中国非制造业商务活动指数为54.4%</title><p>2016/01/01</p>
          建筑业商务活动指数为58.3%，建筑业新订单指数为52.1%。"""
        self.assertEqual(fc.parse_pmi(html, "https://www.stats.gov.cn/sj/zxfb/202302/a.html")[0], "2015-12-31")
        generic = html.replace("<title>12月中国非制造业商务活动指数为54.4%</title>",
                               "<title>国家统计局信息公开</title><h1>2015年12月中国采购经理指数运行情况</h1>")
        self.assertEqual(fc.parse_pmi(generic, "https://www.stats.gov.cn/a.html")[0], "2015-12-31")

    def test_housing_selects_saar_sheet_and_national_total(self):
        w = openpyxl.Workbook()
        w.active.title = "Annual"
        w.active.append([datetime(1959, 1, 1), 99999])
        s = w.create_sheet("Seasonally Adjusted")
        s.append(["(Thousands of units)"])
        s.append(["Seasonally adjusted annual rate"])
        for offset in range(600):
            s.append([datetime(1959 + offset // 12, offset % 12 + 1, 1), 1200 + offset, 999])
        s.append([datetime(2099, 1, 1), 99999])
        pts = fc.parse_housing(workbook_bytes(w), "2026-10-05")
        self.assertEqual(len(pts), 600)
        self.assertEqual(pts[0], ["1959-01-31", 1200])
        self.assertEqual(pts[-1], ["2008-12-31", 1799])

    def test_spending_reordered_columns_revision_flags_and_unit_conversion(self):
        w = openpyxl.Workbook()
        s = w.active
        s.append(["Seasonally Adjusted Annual Rate"])
        s.append(["Millions of dollars"])
        s.append(["Date", "Total\n_x000D_Nonresidential", "Total\n_x000D_Residential"])
        for offset in range(400):
            year, month = 1993 + offset // 12, offset % 12 + 1
            label = f"{fc.calendar.month_abbr[month]}-{year % 100:02}" + ("p" if offset == 399 else "r")
            s.append([label, 1300000, 890000])
        s.append(["Dec-50p", 999999, 999999])
        data = fc.parse_spending(workbook_bytes(w), "2026-10-05")
        self.assertEqual(data["주거용"][0], ("1993-01-31", 890.0))
        self.assertEqual(data["비주거용"][-1], ("2026-04-30", 1300.0))
        self.assertEqual(len(data["주거용"]), 400)

    def test_yoy_uses_calendar_year_not_twelve_available_rows(self):
        pts = [["2024-02-29", 100], ["2024-04-30", 200], ["2025-02-28", 110], ["2025-03-31", 120]]
        self.assertEqual(fc.yoy(pts), [["2025-02-28", 10.0]])

    def test_failure_keeps_data_and_does_not_block_other_country(self):
        with tempfile.TemporaryDirectory() as directory, patch.object(common, "DATA_DIR", Path(directory)):
            p = Path(directory) / "commodities" / "comm_copper_cn_starts.json"
            p.parent.mkdir()
            doc = {"id": "comm_copper_cn_starts", "updated": "2026-08-31", "fetched": "2026-10-01",
                   "series": {"x": [["2000-02-29", 1], ["2026-08-31", 2]]}}
            p.write_text(json.dumps(doc), encoding="utf-8")
            with patch.object(fc, "fetch_cn", side_effect=[ValueError("offline"), None, None, None]), \
                 patch.object(fc, "fetch_pmi"), patch.object(fc, "fetch_us_housing") as housing, \
                 patch.object(fc, "fetch_us_spending") as spending:
                with self.assertRaises(SystemExit):
                    fc.run()
            after = json.loads(p.read_text(encoding="utf-8"))
            self.assertEqual(after["series"], doc["series"])
            self.assertEqual(after["fetched"], doc["fetched"])
            self.assertEqual(after["updated"], doc["updated"])
            self.assertFalse(after["collection_status"]["ok"])
            housing.assert_called_once()
            spending.assert_called_once()


if __name__ == "__main__":
    unittest.main()

"""Copper demand context: official NBS China and US Census construction data.

No API key required. --backfill retrieves construction PMI releases from 2015.
Normal runs merge revisions and check recent PMI releases without deleting history.
"""
import argparse
import calendar
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime
import io
import json
import math
import re
import sys
import time
from urllib.parse import urlparse

import openpyxl
import requests
from bs4 import BeautifulSoup

from common import DATA_DIR, collection_date, load_indicator, merge_points, record_fetch_failure, save_indicator

NBS_PAGE = "https://data.stats.gov.cn/dg/website/page.html#/pc/national/monthData"
NBS_API = "https://data.stats.gov.cn/dg/website/publicrelease/web/external/stream/esData"
NBS_ROOT = "fc982599aa684be7969d7b90b1bd0e84"
SEARCH_API = "https://api.so-gov.cn/query/s"  # Used by stats.gov.cn/search/s.
NRC_PAGE = "https://www.census.gov/construction/nrc/data/series.html"
VIP_PAGE = "https://www.census.gov/construction/c30/historical_data.html"
PREFIX = "comm_copper_"
AREA = "cac0766314e045ea82f69886aabd31b0"
CN_CARDS = [
    ("cn_starts", "중국 부동산 신규 착공", AREA,
     "17660a75e3494e389bad3658ff124995", "d0bfd7e4b56a4bb98cea7cfd141475d9", "백만㎡", 100,
     "새로 착공한 건물의 연초 이후 누적 면적. 향후 건축 물량과 배선·설비 수요를 살피는 지표입니다."),
    ("cn_completions", "중국 부동산 준공", AREA,
     "a0def4ba92f94107b884c46b70adc08b", "526de94de84b4b08a898d94a212d18b8", "백만㎡", 100,
     "완공·인도 가능한 건물의 연초 이후 누적 면적. 구리 배선·설비 설치는 준공 신고보다 앞설 수 있습니다."),
    ("cn_sales", "중국 신축 상품부동산 판매면적", "0ae633cdb85f4a8397650831b2b27e50",
     "d353226cf0434c929b6299f8d4987754", "50a37fbef1d04be68f15d82b711783bf", "백만㎡", 100,
     "신축 주택·상업용 부동산의 계약 면적. 기존 주택 거래 전체를 포함하지 않습니다."),
    ("cn_investment", "중국 부동산 개발투자", "9206137ccf03460daa74b7799e0f3c31",
     "bfb626c0dfa04afab67937c452ca9f50", "205e08cba8c2409980db58c98da91b6f", "십억 위안", 10,
     "연초 이후 누적 개발투자. 명목 금액이며 건축·토지개발·토지 취득 비용 등이 포함됩니다."),
]
PMI_ID = PREFIX + "cn_construction_pmi"
ACTIVITY = "건설 사업활동"
ORDERS = "건설 신규주문"
PMI_CATEGORY = "7a64a6e25aec4a8e9dde044ecd9e2cce"
PMI_ACTIVITY_ID = "97a152f401d64409892a7fdbe22b0225"
MONTHS = {name: i for i, name in enumerate(calendar.month_abbr) if name}


def session():
    s = requests.Session()
    s.headers.update({"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
                      "Referer": "https://data.stats.gov.cn/dg/website/page.html"})
    return s


def request(s, method, url, **kwargs):
    for attempt in range(3):
        try:
            response = s.request(method, url, timeout=(8, 40), **kwargs)
            response.raise_for_status()
            return response
        except requests.RequestException:
            if attempt == 2:
                raise RuntimeError(f"{urlparse(url).hostname} request failed") from None
            time.sleep(attempt + 1)


def month_end(year, month):
    return date(int(year), int(month), calendar.monthrange(int(year), int(month))[1]).isoformat()


def number(value):
    try:
        v = float(str(value).replace(",", "").strip())
        return v if math.isfinite(v) else None
    except (ValueError, TypeError):
        return None


def nbs_series(payload, ids, today=None):
    """Only declared IDs and valid observed monthly periods; empty is never zero."""
    if payload.get("success") is not True or not isinstance(payload.get("data"), list):
        raise ValueError("NBS returned an unsuccessful or malformed response")
    cutoff = today or collection_date()
    points = {i: {} for i in ids}
    for row in payload["data"]:
        match = re.fullmatch(r"(\d{4})(\d{2})MM", str(row.get("code", "")))
        if not match:
            continue
        dt = month_end(*match.groups())
        if dt > cutoff:
            continue
        for item in row.get("values", []):
            identifier = item.get("_id")
            v = number(item.get("value"))
            if identifier in points and v is not None:
                points[identifier][dt] = v
    result = {i: sorted(v.items()) for i, v in points.items()}
    if any(not values for values in result.values()):
        raise ValueError("NBS requested series is empty; preserve existing history")
    return result


def query_nbs(s, category, ids):
    today = collection_date()
    payload = {"cid": category, "rootId": NBS_ROOT, "indicatorIds": ids,
               "daCatalogId": "", "das": [{"text": "全国", "value": "000000000000"}],
               "showType": 1, "dts": [f"200001MM-{today[:7].replace('-', '')}MM"]}
    return nbs_series(request(s, "POST", NBS_API, json=payload).json(), ids)


def merged(existing, incoming):
    values = dict(existing or [])
    values.update(dict(incoming))
    return [[dt, values[dt]] for dt in sorted(values)]


def base_doc(identifier, name, source, url):
    return {"id": identifier, "name": name, "frequency": "monthly", "source": source,
            "source_url": url, "month_labels": True, "monthly_axis": True,
            "data_stale_days": 100, "change_mode": "none", "span_gaps": False, "series": {}}


def cn_scale_views(level, unit):
    """Derived quantities never sum overlapping YTD observations."""
    values = dict(level)
    monthly, annotations = {}, {}
    for dt, value in sorted(values.items()):
        year, month = int(dt[:4]), int(dt[5:7])
        if month == 1:
            monthly[dt] = value
        elif month == 2 and month_end(year, 1) not in values:
            for m in [1, 2]:
                date_key = month_end(year, m)
                monthly[date_key] = round(value / 2, 4)
                annotations[date_key] = "추정 · 1~2월 누적 물량을 절반씩 배분"
        else:
            prior = values.get(month_end(year, month - 1))
            # A missing prior month or a negative revision cannot be monthly output.
            if prior is not None and value >= prior:
                monthly[dt] = round(value - prior, 4)
    complete_years = [int(dt[:4]) for dt in values if dt[5:7] == "12"]
    total, start = [], min(complete_years) if complete_years else None
    if start is not None:
        carried = 0
        for year in range(start, max(int(dt[:4]) for dt in values) + 1):
            total.extend([[dt, round(carried + value, 4)] for dt, value in sorted(values.items())
                          if int(dt[:4]) == year])
            annual = values.get(month_end(year, 12))
            if annual is None:
                break  # Do not silently resume across a missing annual total.
            carried += annual
    views = {
        "monthly": {"label": "월별 막대", "unit": unit, "series": {"월별 규모 (계산)": sorted(monthly.items())},
                    "description": "공식 연초 누적 자료에서 계산한 월별 물량입니다. 월별 증감과 계절적 흐름을 비교할 수 있습니다.",
                    "default_series": ["월별 규모 (계산)"], "cumulative": False,
                    "bridge_missing_january": False, "chart_type": "bar", "zero_baseline": True,
                    "point_annotations": annotations,
                    "note": "누적 규모의 전월 차이로 계산합니다. 1~2월은 절반씩 배분한 추정치이며 실제 월별 실적과 다를 수 있습니다. 원자료 수정으로 전월 차이가 음수이거나 전월 자료가 없으면 비워둡니다."},
        "total": {"label": "전체 누적", "unit": unit, "series": {"전체 누적 규모": total},
                  "description": "각 연도의 실적을 이어 더한 장기 합계입니다. 현재 잔존 건물이나 재고 규모를 뜻하지 않습니다.",
                  "default_series": ["전체 누적 규모"], "cumulative": False, "cumulative_since": start,
                  "bridge_missing_january": True,
                  "note": f"{start}년부터의 이전 연도 연간 실적 + 현재 연초 누적입니다. 연초 누적 값을 매월 중복 합산하지 않습니다. 1월 미발표 구간은 점선으로 연결합니다. 통계 범위·수정 기준이 바뀌어 장기 합계는 참고용입니다."},
    }
    return {"monthly": views["monthly"],
            "monthly_line": {**views["monthly"], "label": "월별 선그래프", "chart_type": "line"},
            "total": views["total"]}


def set_cn_views(doc, level, unit):
    label = "누적 전년동기비"
    doc.update(bridge_missing_january=True,
               note="1월은 별도 발표하지 않으며 2월은 1~2월 합산입니다. 12월~2월의 점선은 흐름 연결이며 1월 추정값이 아닙니다. 증가율은 NBS의 공식 비교 가능 기준입니다.")
    doc["series_views"] = {
        "yoy": {"label": "누적 증가율", "unit": "%", "series": doc["series"], "default_series": [label]},
        "level": {"label": "누적 규모", "unit": unit, "series": {"누적 규모": level}, "default_series": ["누적 규모"],
                  "bridge_missing_january": False,
                  "note": "연초 이후 누적 규모입니다. 1월은 별도 발표하지 않으며 2월은 1~2월 합산입니다."},
    }
    doc["series_views"].update(cn_scale_views(level, unit))


def fetch_cn(s, card):
    suffix, name, category, level_id, growth_id, unit, divisor, description = card
    identifier = PREFIX + suffix
    result = query_nbs(s, category, [level_id, growth_id])
    old = load_indicator("commodities", identifier)
    label = "누적 전년동기비"
    doc = base_doc(identifier, name, "중국 국가통계국 (NBS)", NBS_PAGE)
    doc.update(unit="%", default_series=[label], cumulative=True, description=description,
               bridge_missing_january=True,
               note="1월은 별도 발표하지 않으며 2월은 1~2월 합산입니다. 12월~2월의 점선은 흐름 연결이며 1월 추정값이 아닙니다. 증가율은 NBS의 공식 비교 가능 기준입니다.")
    doc["series"][label] = merged(old.get("series", {}).get(label), result[growth_id])
    old_level = old.get("series_views", {}).get("level", {}).get("series", {}).get("누적 규모", [])
    level = merged(old_level, [[dt, round(v / divisor, 4)] for dt, v in result[level_id]])
    set_cn_views(doc, level, unit)
    doc["default_view"] = "yoy"
    doc["nbs_identifiers"] = {"category": category, "level": level_id, "growth": growth_id}
    save_indicator("commodities", doc, data_date=True)


def parse_pmi(html, url):
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text() if soup.title else ""
    # Modern combined releases and older separate non-manufacturing releases.
    text = re.sub(r"\s+", "", soup.get_text())
    # The information-disclosure mirror has a generic HTML title but a dated heading.
    heading = title + " " + text[:1200]
    match = re.search(r"(20\d{2})年(\d{1,2})月中国采购经理指数运行情况", heading)
    published = re.search(r"(20\d{2})/(\d{2})/(\d{2})", text)
    if match:
        year, month = map(int, match.groups())
    else:
        match = re.search(r"(?:(20\d{2})年)?(\d{1,2})月中国非制造业商务活动指数", heading)
        if not match or not published:
            raise ValueError("Not a monthly NBS construction PMI release")
        month = int(match.group(2))
        year = int(match.group(1)) if match.group(1) else int(published.group(1)) - (1 if int(published.group(2)) == 1 and month == 12 else 0)
    dt = month_end(year, month)
    values = {}
    for label, phrase in [(ACTIVITY, "建筑业商务活动指数"), (ORDERS, "建筑业新订单指数")]:
        hits = re.findall(phrase + r"为([\d.]+)[%％]", text)
        unique = {float(v) for v in hits}
        if len(unique) != 1 or not 0 <= next(iter(unique), -1) <= 100:
            raise ValueError(f"Ambiguous or missing construction component: {label}")
        values[label] = next(iter(unique))
    if dt > collection_date():
        raise ValueError("Future observation rejected")
    publication = "-".join(published.groups()) if published else None
    return dt, values, {"url": url, "label": f"NBS {year}-{month:02}", "published": publication}


def pmi_links(s, backfill=False):
    params = {"siteCode": "bm36000002", "tab": "all", "sort": "dateDesc", "adv": 1,
              "qt": '("建筑业新订单指数")', "keyPlace": 0, "timeOption": 2,
              "startDateStr": "2015-01-01", "endDateStr": collection_date()}
    links = {}
    for page in range(1, 31 if backfill else 4):
        response = request(s, "POST", SEARCH_API, data={**params, "page": page}).json()
        if not response.get("ok"):
            raise ValueError("NBS release search failed")
        docs = response.get("resultDocs", [])
        for item in docs:
            data = item.get("data", {})
            url = data.get("url", "")
            title = data.get("titleO", "")
            if urlparse(url).hostname != "www.stats.gov.cn":
                continue
            if not (re.fullmatch(r"20\d{2}年\d{1,2}月中国采购经理指数运行情况", title)
                    or re.match(r"(?:20\d{2}年)?\d{1,2}月中国非制造业商务活动指数", title)):
                continue
            links[title] = url.replace("http://", "https://", 1)
        if page * 20 >= response.get("totalHits", 0) or not docs:
            break
    return list(links.values())


def fetch_pmi(backfill=False):
    s = session()
    old = load_indicator("commodities", PMI_ID)
    doc = base_doc(PMI_ID, "중국 건설 경기 · 사업활동 / 신규주문", "중국 국가통계국 · CFLP", "https://www.stats.gov.cn/sj/zxfb/")
    doc.update(unit="pt", default_series=[ACTIVITY, ORDERS], reference_value=50,
               reference_label="확장·위축 기준 50", change_mode="percentage_points",
               description="건설 기업의 계절조정 설문 지수. 50 초과는 전월보다 확장, 50 미만은 위축을 뜻합니다.",
               note="사업활동은 당월 경기, 신규주문은 후속 일감을 살피는 지표입니다. 구리 소비량이나 증가율을 직접 측정하지 않습니다.")
    doc["series"] = old.get("series", {})
    doc["point_sources"] = old.get("point_sources", {})
    warnings = []
    try:
        activity = query_nbs(s, PMI_CATEGORY, [PMI_ACTIVITY_ID])[PMI_ACTIVITY_ID]
        merge_points(doc, ACTIVITY, activity)
    except Exception as e:
        warnings.append(str(e))
    links = pmi_links(s, backfill)
    def fetch_one(url):
        response = request(session(), "GET", url)
        response.encoding = "utf-8"
        return parse_pmi(response.text, url)
    records = []
    with ThreadPoolExecutor(max_workers=4) as pool:
        futures = {pool.submit(fetch_one, url): url for url in links}
        for future in as_completed(futures):
            try:
                records.append(future.result())
            except Exception:
                warnings.append("일부 과거 발표문 확인 실패")
    # Use canonical new release URLs consistently when duplicate releases exist.
    for dt, values, source in sorted(records, key=lambda row: (row[0], row[2]["url"])):
        for label, value in values.items():
            merge_points(doc, label, [[dt, value]])
        doc["point_sources"][dt] = source
    if not records or not all(doc["series"].get(label) for label in [ACTIVITY, ORDERS]):
        raise ValueError("No verified construction PMI release; preserve existing document")
    latest = max(dt for dt, _, _ in records)
    for label in [ACTIVITY, ORDERS]:
        if not any(dt == latest for dt, _ in doc["series"][label]):
            raise ValueError("Latest PMI components are not aligned")
    doc["published"] = doc["point_sources"][latest].get("published")
    if warnings:
        doc["collection_status"] = {"ok": False, "checked": collection_date(), "kind": "partial",
                                    "message": "최신 발표 확인 · 일부 과거 자료 확인 실패, 기존 이력 보존"}
    save_indicator("commodities", doc, data_date=True)
    print(f"  PMI: {len(records)} verified releases")


def parse_housing(content, today=None, seasonally_adjusted=True):
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook["Seasonally Adjusted" if seasonally_adjusted else "Not Seasonally Adjusted"]
        header = " ".join(str(cell) for row in list(sheet.values)[:6] for cell in row if cell is not None)
        basis = "Seasonally adjusted annual rate" if seasonally_adjusted else "Not seasonally adjusted"
        if basis not in header or "Thousands of units" not in header:
            raise ValueError("Census housing units/seasonal basis changed")
        points = []
        for row in sheet.values:
            if isinstance(row[0], (datetime, date)) and number(row[1]) is not None:
                dt = month_end(row[0].year, row[0].month)
                if dt <= (today or collection_date()):
                    points.append([dt, number(row[1])])
        if len(points) < 600:
            raise ValueError("Census housing history unexpectedly short")
        return sorted(points)
    finally:
        workbook.close()


def parse_spending(content, today=None):
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        rows = list(workbook.active.values)
        header = " ".join(str(v) for row in rows[:5] for v in row if v is not None)
        if "Seasonally Adjusted Annual Rate" not in header or "Millions of dollars" not in header:
            raise ValueError("Census spending units/seasonal basis changed")
        columns = next((row for row in rows if row[0] == "Date"), None)
        normalized = [re.sub(r"\s+", " ", str(v).replace("_x000D_", " ")).strip() for v in columns or []]
        indices = {label: normalized.index(column) for label, column in
                   [("주거용", "Total Residential"), ("비주거용", "Total Nonresidential")]}
        result = {label: {} for label in indices}
        for row in rows:
            match = re.fullmatch(r"([A-Za-z]{3})-(\d{2})([pr]?)", str(row[0]).strip())
            if not match:
                continue
            year2 = int(match.group(2))
            year = 1900 + year2 if year2 >= 93 else 2000 + year2
            dt = month_end(year, MONTHS[match.group(1).title()])
            if dt > (today or collection_date()):
                continue
            for label, col in indices.items():
                v = number(row[col])
                # Unavailable pre-2002 category totals are encoded as 0.
                # National residential/nonresidential spending cannot be zero.
                if v is not None and v > 0:
                    result[label][dt] = round(v / 1000, 3)  # USD millions -> billions.
        if any(len(values) < 250 for values in result.values()):
            raise ValueError("Census spending history unexpectedly short")
        return {label: sorted(values.items()) for label, values in result.items()}
    finally:
        workbook.close()


def yoy(points):
    values = dict(points)
    out = []
    for dt, value in sorted(values.items()):
        prior = month_end(int(dt[:4]) - 1, int(dt[5:7]))
        base = values.get(prior)
        if base is not None and base > 0:
            out.append([dt, round((value / base - 1) * 100, 3)])
    return out


def cumulative_months(points, yearly=False, start=None):
    """Sum monthly counts; missing months never become zero or partial totals."""
    values = dict(points)
    if not values:
        return []
    offset = lambda dt: int(dt[:4]) * 12 + int(dt[5:7]) - 1
    first, last = offset(start or min(values)), offset(max(values))
    total, valid, result = 0, True, []
    for month in range(first, last + 1):
        year, m = divmod(month, 12)
        if yearly and (month == first or m == 0):
            total, valid = 0, m == 0
        dt = month_end(year, m + 1)
        if dt not in values:
            if not yearly:
                break
            valid = False
        if valid:
            total += values[dt]
            result.append([dt, round(total, 3)])
    return result


def set_us_housing_cumulative_views(doc):
    counts = doc["monthly_counts"]
    start = max(min(dt for dt, _ in pts) for pts in counts.values())
    common = {"unit": "천 호", "default_series": list(counts), "chart_type": "line",
              "zero_baseline": True,
              "description": "Census의 비계절조정 월별 주택 물량을 합산한 값입니다. 착공·허가·준공은 별개 지표이며 서로 더하지 않습니다."}
    doc["series_views"]["ytd"] = {**common, "label": "연초 누적", "cumulative": True,
        "series": {label: cumulative_months(pts, yearly=True) for label, pts in counts.items()},
        "note": "매년 1월부터 해당 월까지의 물량 합계입니다. 계절성이 포함되며 연율(SAAR)을 합산하거나 나눈 추정치가 아닙니다."}
    doc["series_views"]["total"] = {**common, "label": "전체 누적", "cumulative_since": int(start[:4]), "zero_baseline": False,
        "series": {label: cumulative_months(pts, start=start) for label, pts in counts.items()},
        "note": f"세 지표의 공통 관측 시작인 {start[:7]}부터 월별 물량을 합산합니다. 조회 기간을 바꿔도 누적 기준은 유지됩니다. 현재 잔존 주택이나 미완공 재고를 뜻하지 않습니다."}


def save_us(identifier, name, unit, incoming, url, description, old=None, monthly_counts=None):
    old = old or load_indicator("commodities", identifier)
    doc = base_doc(identifier, name, "미국 Census Bureau", url)
    doc.update(unit=unit, default_series=list(incoming), description=description,
               note="계절조정 연율(SAAR): 해당 월의 속도를 연간으로 환산한 값입니다. 당월 실적이나 연간 누적이 아닙니다.")
    doc["series"] = {label: merged(old.get("series", {}).get(label), pts) for label, pts in incoming.items()}
    doc["series_views"] = {
        "level": {"label": "규모 (연율)", "unit": unit, "series": doc["series"], "default_series": list(incoming)},
        "yoy": {"label": "전년동월비", "unit": "%", "series": {label: yoy(pts) for label, pts in doc["series"].items()}, "default_series": list(incoming)},
    }
    if monthly_counts is not None:
        doc["monthly_counts"] = {label: merged(old.get("monthly_counts", {}).get(label), pts)
                                 for label, pts in monthly_counts.items()}
        set_us_housing_cumulative_views(doc)
    doc["default_view"] = "level"
    save_indicator("commodities", doc, data_date=True)


def fetch_us_housing(s):
    identifier = PREFIX + "us_housing"
    old = load_indicator("commodities", identifier)
    incoming, monthly_counts, failures = {}, {}, []
    files = [("주택 착공", "starts"), ("건축허가", "permits"), ("주택 준공", "comps")]
    for label, filename in files:
        try:
            url = f"https://www.census.gov/construction/nrc/xls/{filename}_cust.xlsx"
            content = request(s, "GET", url).content
            annualized = parse_housing(content)
            monthly = parse_housing(content, seasonally_adjusted=False)
            incoming[label], monthly_counts[label] = annualized, monthly
        except Exception as e:
            failures.append(label)
            if not old.get("series", {}).get(label) or not old.get("monthly_counts", {}).get(label):
                raise e
            incoming[label] = old["series"][label]
            monthly_counts[label] = old["monthly_counts"][label]
    if len(failures) == len(files):
        raise ValueError("All Census housing sources failed")
    save_us(identifier, "미국 주택 건설 · 착공 / 허가 / 준공", "천 호/년", incoming, NRC_PAGE,
            "전국 민간 신규 주택의 건축허가·착공·준공 물량. 기존 주택 매매는 포함하지 않습니다.", old, monthly_counts)
    if failures:
        record_fetch_failure("commodities", identifier, ValueError("일부 항목 수집 실패: " + ", ".join(failures)))


def fetch_us_spending(s):
    url = "https://www.census.gov/construction/c30/xlsx/totsatime.xlsx"
    incoming = parse_spending(request(s, "GET", url).content)
    identifier = PREFIX + "us_construction_spending"
    old = load_indicator("commodities", identifier)
    # Remove previously imported placeholders, preserving real history.
    old["series"] = {label: [[dt, value] for dt, value in pts if value > 0]
                     for label, pts in old.get("series", {}).items()}
    save_us(identifier, "미국 건설 지출 · 주거용 / 비주거용", "십억 달러/년", incoming,
            VIP_PAGE, "민간·공공의 주거용 및 비주거용 공사 지출. 신축과 기존 시설 개선을 포함하며 명목 금액입니다. 세부 항목은 2002년부터 제공됩니다.", old)


def run(backfill=False, country="all"):
    jobs = []
    if country in ["all", "cn"]:
        jobs.extend((PREFIX + card[0], lambda card=card: fetch_cn(session(), card)) for card in CN_CARDS)
        jobs.append((PMI_ID, lambda: fetch_pmi(backfill)))
    if country in ["all", "us"]:
        jobs.extend([(PREFIX + "us_housing", lambda: fetch_us_housing(session())),
                     (PREFIX + "us_construction_spending", lambda: fetch_us_spending(session()))])
    errors = []
    for identifier, job in jobs:
        try:
            job()
        except Exception as e:
            record_fetch_failure("commodities", identifier, e)
            errors.append(identifier)
            print(f"  FAIL {identifier}: {type(e).__name__}: {e}", file=sys.stderr)
    if errors:
        raise SystemExit("Construction sources failed: " + ", ".join(errors))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--backfill", action="store_true")
    parser.add_argument("--country", choices=["all", "cn", "us"], default="all")
    args = parser.parse_args()
    run(args.backfill, args.country)

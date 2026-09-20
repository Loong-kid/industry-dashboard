# -*- coding: utf-8 -*-
"""미국 LNG 액화 프로젝트 → data/lng/*.json (표 + 용량 시계열).

원본: EIA 'U.S. Liquefaction Capacity' 워크북(xlsx, 키 없음).
  https://www.eia.gov/naturalgas/importsexports/liquefactioncapacity/U.S.liquefactioncapacity.xlsx
  시트 2개 — 'Existing & Under Construction'(FID 완료: 가동·시운전·건설중, train 단위)
            'Approved'(FID 전: 허가는 났고 FEED 단계, 프로젝트 단위)

**이 워크북은 분기 갱신이라고 적혀 있지만 실제로는 2025-06-30 이후 멈춰 있다**(2026-09 확인).
그래서 릴리스 날짜를 카드에 노출하고, 빈티지를 저장해 다음 발표가 나오면 무엇이 바뀌었는지
(가동 예정연도가 밀렸는지) 비교할 수 있게 해 둔다 — 전력 탭 EIA-860M과 같은 방식이다.

    python scripts/fetch_lng_projects.py
"""
import csv
import datetime as dt
import io
import json
import re
from pathlib import Path

import openpyxl
import requests

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "data" / "lng"          # 탭 id=lng
VINTAGE_DIR = ROOT / "data" / "_lng"     # 발표분 보관(지연 추적용)
SRC = "https://www.eia.gov/naturalgas/importsexports/liquefactioncapacity/U.S.liquefactioncapacity.xlsx"

# EIA 상태 문자열 → 화면 표기. 워크북이 각주 문자를 붙여 쓰는 경우가 있어(예: 'Under constructionH')
# 앞부분만 보고 맞춘다.
STATUS_MAP = [
    ("commercial operation", "가동"),
    ("commissioning", "시운전"),
    ("under construction", "건설중"),
]


def clean_status(raw):
    s = str(raw or "").strip().lower()
    for key, label in STATUS_MAP:
        if s.startswith(key):
            return label
    return None


def num(v):
    """'6.75^', '5.520', 수식 잔재('=C8*E8') 등이 섞여 있다 → 숫자만 뽑는다."""
    if isinstance(v, (int, float)):
        return round(float(v), 3)
    m = re.search(r"-?\d+(?:\.\d+)?", str(v or ""))
    return round(float(m.group()), 3) if m else None


def year_of(v):
    """가동 시점은 날짜('2025-09-01')일 때도, 반기 표기('2H2026')나 연도('2027')일 때도 있다."""
    if isinstance(v, dt.datetime):
        return v.year
    m = re.search(r"(19|20)\d{2}", str(v or ""))
    return int(m.group()) if m else None


def fetch_workbook():
    r = requests.get(SRC, headers={"User-Agent": "Mozilla/5.0"}, timeout=(5, 60))
    r.raise_for_status()
    wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
    release = None
    for row in wb["Contents"].iter_rows(values_only=True):
        for i, c in enumerate(row):
            if str(c or "").strip().lower().startswith("release date") and i + 1 < len(row):
                release = year_month_day(row[i + 1])
    return wb, release or dt.date.today().isoformat()


def year_month_day(v):
    return v.date().isoformat() if isinstance(v, dt.datetime) else (str(v)[:10] if v else None)


def parse_built(ws):
    """FID 완료 시트: train 단위. 각주 문단이 표 아래 붙어 있어 상태가 없는 행은 버린다."""
    rows = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        status = clean_status(r[6])
        if not r[0] or not status:
            continue
        rows.append({
            "project": str(r[0]).strip(), "train": str(r[1] or "").strip(),
            "mtpa": num(r[3]), "status": status, "year": year_of(r[7]),
            "state": str(r[9] or "").strip() if len(r) > 9 else "",
            "phase": "FID 완료",
        })
    return rows


def parse_approved(ws):
    """FID 전 시트: 프로젝트 단위. 용량은 '전체 설계용량'(트레인 수 반영) 컬럼을 쓴다."""
    rows = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        name = str(r[0] or "").strip()
        if not name or name.lower().startswith(("notes", "^", "source")) or len(name) > 60:
            continue
        mtpa = num(r[6]) or num(r[3])
        if mtpa is None:
            continue
        rows.append({
            "project": name, "train": f"{num(r[4]) or 1:.0f}개 트레인" if r[4] else "",
            "mtpa": mtpa, "status": "승인(FID 전)", "year": None,
            "state": str(r[8] or "").strip() if len(r) > 8 else "",
            "operator": str(r[1] or "").strip(),
            "feed": str(r[7] or "").strip(),
            "phase": "FID 전",
        })
    return rows


def save(cid, doc):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc["fetched"] = dt.date.today().isoformat()
    (OUT_DIR / f"{cid}.json").write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")


def build_timeline(built, approved, release):
    """연도별 신규 가동 용량. 이미 가동한 물량과 앞으로 들어올 물량을 갈라서 보여준다 —
    합쳐 놓으면 '계획이 실적처럼' 보인다(전력 탭에서 얻은 교훈)."""
    this_year = dt.date.today().year
    cats = {"가동 개시(실적)": {}, "건설중(예정)": {}, "승인·FID 전(시기 미정)": {}}
    for r in built:
        if not r["year"] or not r["mtpa"]:
            continue
        key = "가동 개시(실적)" if r["status"] == "가동" and r["year"] <= this_year else "건설중(예정)"
        cats[key][r["year"]] = round(cats[key].get(r["year"], 0) + r["mtpa"], 2)
    pending = round(sum(r["mtpa"] for r in approved if r["mtpa"]), 2)

    series = {k: [[f"{y}-01-01", v] for y, v in sorted(d.items())] for k, d in cats.items() if d}
    doc = {
        "id": "lng_capacity_timeline", "name": "미국 LNG 액화용량 — 연도별 신규",
        "unit": "MTPA", "frequency": "quarterly",
        "source": f"EIA U.S. Liquefaction Capacity ({release} 발표)",
        "source_url": SRC, "updated": release,
        "default_series": [k for k in series],
        "description": (
            "미국 액화 트레인이 해마다 얼마나 새로 돌기 시작했고, 앞으로 얼마가 들어올지를 나눠 담았다. "
            "MTPA는 연간 백만 톤 기준 설비 용량이다. '실적'은 이미 상업운전에 들어간 물량, '건설중'은 "
            "FID(최종투자결정)를 마치고 짓는 중이라 가동 시점이 정해진 물량이다. 둘을 합쳐 보면 계획이 "
            "실적처럼 보이므로 일부러 갈라 두었다."
        ),
        "note": (f"FID 전 승인 물량은 가동 시기가 정해지지 않아 시계열에 넣지 않았다 — 합계 {pending:,.1f} MTPA는 "
                 f"아래 표에서 본다. 원본은 분기 갱신을 표방하지만 {release} 이후 새 발표가 없다."),
        "series": series,
    }
    save("lng_capacity_timeline", doc)

    # 누적: 지금까지 깔린 총 용량 + 건설중까지 다 들어왔을 때의 경로
    cum, total = [], 0.0
    for y, v in sorted({**cats["가동 개시(실적)"]}.items()):
        total = round(total + v, 2)
        cum.append([f"{y}-01-01", total])
    fut, ftotal = [], total
    for y, v in sorted(cats["건설중(예정)"].items()):
        ftotal = round(ftotal + v, 2)
        fut.append([f"{y}-01-01", ftotal])
    save("lng_capacity_cumulative", {
        "id": "lng_capacity_cumulative", "name": "미국 LNG 액화용량 — 누적",
        "unit": "MTPA", "frequency": "quarterly",
        "source": f"EIA U.S. Liquefaction Capacity ({release} 발표)",
        "source_url": SRC, "updated": release,
        "default_series": ["가동중 누적", "건설중 포함 전망"],
        "description": (
            "가동에 들어간 용량을 계속 더한 값과, 지금 짓고 있는 물량까지 다 들어왔을 때의 경로. "
            "미국이 LNG 수출국으로 올라선 속도와 앞으로 몇 년간 더해질 크기를 한 선에서 본다."
        ),
        "note": "전망선은 EIA가 적어 둔 가동 예정연도를 그대로 쌓은 것이라 지연되면 오른쪽으로 밀린다.",
        "series": {"가동중 누적": cum, "건설중 포함 전망": fut},
    })
    return pending


def build_table(built, approved, release, pending):
    rows = []
    for r in built + approved:
        rows.append({
            "project": r["project"], "train": r["train"], "mtpa": r["mtpa"],
            "status": r["status"], "year": r["year"], "state": r["state"],
            "detail": r.get("feed", ""),
        })
    save("lng_projects", {
        "id": "lng_projects", "name": f"미국 LNG 프로젝트 ({release} 발표 기준)",
        "source": "EIA U.S. Liquefaction Capacity", "source_url": SRC, "updated": release,
        "note": (f"FID를 마친 트레인 {len(built)}건과 허가만 난 FID 전 프로젝트 {len(approved)}건"
                 f"(합계 {pending:,.1f} MTPA). 상태·용량 컬럼을 눌러 정렬한다. "
                 f"FID 전 물량은 착공 전이라 상당수가 지연되거나 취소된다."),
        "cols": [
            {"key": "project", "label": "프로젝트"},
            {"key": "train", "label": "트레인"},
            {"key": "mtpa", "label": "용량(MTPA)", "align": "right", "fmt": "num"},
            {"key": "status", "label": "상태"},
            {"key": "year", "label": "가동(예정)", "align": "right"},
            {"key": "state", "label": "주"},
            {"key": "detail", "label": "진행 단계"},
        ],
        "rows": rows,
    })


def save_vintage(built, approved, release):
    """발표분을 그대로 보관한다. 다음 발표가 나오면 가동 예정연도가 어떻게 움직였는지 비교해
    '지연'을 잡아낼 수 있다 — 지금은 빈티지가 하나뿐이라 비교는 다음 발표부터."""
    VINTAGE_DIR.mkdir(parents=True, exist_ok=True)
    path = VINTAGE_DIR / f"eia_{release}.csv"
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["project", "train", "mtpa", "status", "year", "state", "phase"])
        w.writeheader()
        for r in built + approved:
            w.writerow({k: r.get(k, "") for k in w.fieldnames})
    return sorted(VINTAGE_DIR.glob("eia_*.csv"))


def build_slip(vintages, release):
    """빈티지 2개 이상이면 프로젝트별 가동 예정연도 변화(+면 지연)를 표로 만든다."""
    if len(vintages) < 2:
        print(f"  빈티지 {len(vintages)}개 — 지연 비교는 다음 발표분부터")
        return
    def load(p):
        with p.open(encoding="utf-8") as f:
            return {f"{r['project']}|{r['train']}": r for r in csv.DictReader(f)}
    old, new = load(vintages[-2]), load(vintages[-1])
    rows = []
    for k, n in new.items():
        o = old.get(k)
        if not o or not n["year"] or not o["year"]:
            continue
        slip = int(n["year"]) - int(o["year"])
        if slip:
            rows.append({"project": n["project"], "train": n["train"], "mtpa": num(n["mtpa"]),
                         "was": int(o["year"]), "now": int(n["year"]), "slip": slip})
    rows.sort(key=lambda r: -r["slip"])
    save("lng_slip", {
        "id": "lng_slip", "name": "가동 예정연도 변화 (발표분 간)",
        "source": "EIA U.S. Liquefaction Capacity", "source_url": SRC, "updated": release,
        "note": (f"{vintages[-2].stem.replace('eia_', '')} → {vintages[-1].stem.replace('eia_', '')} 발표 사이에 "
                 f"가동 예정연도가 바뀐 트레인 {len(rows)}건. 양수면 지연이다."),
        "cols": [
            {"key": "project", "label": "프로젝트"},
            {"key": "train", "label": "트레인"},
            {"key": "mtpa", "label": "용량(MTPA)", "align": "right", "fmt": "num"},
            {"key": "was", "label": "이전 발표", "align": "right"},
            {"key": "now", "label": "현재 발표", "align": "right"},
            {"key": "slip", "label": "변화(년)", "align": "right", "fmt": "int"},
        ],
        "rows": rows,
    })


def run():
    wb, release = fetch_workbook()
    built = parse_built(wb["Existing & Under Construction"])
    approved = parse_approved(wb["Approved"])
    print(f"  EIA 발표 {release}: FID 완료 {len(built)}건 · FID 전 {len(approved)}건")
    pending = build_timeline(built, approved, release)
    build_table(built, approved, release, pending)
    vintages = save_vintage(built, approved, release)
    build_slip(vintages, release)
    print(f"  빈티지 보관: {len(vintages)}개")


if __name__ == "__main__":
    run()

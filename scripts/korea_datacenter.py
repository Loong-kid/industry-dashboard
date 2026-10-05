"""Curated Korean data-centre history + public BLCM permit/start observations.

Keep manual facts and uncertain dates intact. Count permit/construction phases,
never Excel building rows; this is a tracked sample, not a national census.
"""
import copy
import hashlib
import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from common import collection_date

ROOT = Path(__file__).resolve().parent.parent
PUBLIC_URL = 'https://blcm.go.kr/stat/customizedStatic/CustomizedStaticSupplyList.do'
BASELINE = ROOT / 'manual/korea_datacenter_baseline.json'
RESEARCH = ROOT / 'manual/korea_datacenter_research.json'
OUTPUT = ROOT / 'data/datacenter'
PROVINCES = {'서울':'서울특별시','서울시':'서울특별시','경기':'경기도','인천':'인천광역시','인천시':'인천광역시',
    '부산':'부산광역시','부산시':'부산광역시','울산':'울산광역시','울산시':'울산광역시',
    '대전':'대전광역시','대전시':'대전광역시','대구':'대구광역시','광주':'광주광역시',
    '세종':'세종특별자치시','세종시':'세종특별자치시','충남':'충청남도','충북':'충청북도',
    '경남':'경상남도','경북':'경상북도','전남':'전라남도','전북':'전라북도','강원':'강원도'}


def compact(value):
    return re.sub(r'\s+', '', str(value or '')).casefold()


def stable_id(value, prefix='dc'):
    return prefix + '-' + hashlib.sha256(value.encode('utf8')).hexdigest()[:16]


def text(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value):
    try:
        s = text(value).replace(',', '')
        return float(s) if re.fullmatch(r'\d+(?:\.\d+)?', s) else None
    except (TypeError, ValueError):
        return None


def date_info(value, today=None):
    today = today or collection_date()
    s = text(value)
    if not s:
        return {'text': '', 'date': None, 'kind': 'missing', 'precision': None}
    planned = any(k in s for k in ['예정', '목표', '전망'])
    uncertain = '?' in s
    digits = s.replace('.0', '')
    exact = None
    if re.fullmatch(r'\d{8}', digits):
        digits = f'{digits[:4]}-{digits[4:6]}-{digits[6:]}'
    m = re.match(r'^(\d{4})[-./](\d{1,2})[-./](\d{1,2})(?:\s|\(|$)', digits)
    if m:
        try:
            exact = date(*map(int, m.groups())).isoformat()
        except ValueError:
            pass
    year = re.search(r'(?:19|20)\d{2}', s)
    return {'text': s, 'date': exact,
            'kind': 'planned' if planned or (exact and exact > today) else 'confirmed' if exact and not uncertain else 'uncertain',
            'precision': 'day' if exact else 'quarter' if re.search(r'[1-4]Q', s, re.I) else 'year' if year else None,
            'year': int(year.group()) if year else None}


def address_key(location):
    """Municipality + legal-dong + lot; missing lots cannot match automatically."""
    s = compact(location)
    for old, new in sorted(PROVINCES.items(), key=lambda x: -len(x[0])):
        if s.startswith(old) and not s.startswith(new):
            rest = s[len(old):]
            if not rest.startswith(('도','광역시','특별시','특별자치')):
                s = new + rest
            break
    for old, new in [('강원특별자치도','강원도'), ('전북특별자치도','전라북도'),
                     ('서울특별시서울시','서울특별시'), ('경기도경기','경기도')]:
        s = s.replace(old, new)
    # Keep province and municipality, ignore spacing/leading zeroes in lot numbers.
    m = re.search(r'(산)?(\d+)(?:-(\d+))?$', s)
    provinces = set(PROVINCES.values()) | {'전남광주통합특별시'}
    if not m or int(m[2]) == 0 or not any(s.startswith(p) for p in provinces):
        return None
    return s[:m.start()] + ('산' if m[1] else '') + str(int(m[2])) + '-' + str(int(m[3] or 0))


def lot(value):
    n = number(value)
    return str(int(n)) if n is not None else text(value)


def row_location(province, county, dong, bun, ji, land=''):
    province, county = text(province), text(county)
    province = PROVINCES.get(province, province)
    county = re.sub(r'^(?:서울시|서울|경기|인천시|인천|부산시|부산|대구시|대구|대전시|대전|광주시|광주|울산시|울산|전남|전북|경남|경북|충남|충북)(?=.{1,}시|.{1,}구|.{1,}군)', '', county)
    if county == province:
        county = ''
    prefix = ' '.join([province, county, text(dong)])
    return f"{prefix} {'산' if land == '산' else ''}{lot(bun)}-{lot(ji) or '0'}".strip()


def construction_type(value, permit_kind=''):
    s = text(value) or text(permit_kind)
    return next((v for v in ['신축','증축','대수선','용도변경','개축','재축'] if v in s), '미분류')


def parse_public_rows(headers, rows, source, *, starts=False, curated=False):
    """Filter centres, aggregate repeated building rows without repeated total area."""
    idx = {text(h): i for i, h in enumerate(headers) if h is not None}
    if not any(k in idx for k in ['건축물명칭', '건물명']):
        raise ValueError('Missing building-name column')
    def value(row, *names):
        for name in names:
            if name in idx and idx[name] < len(row) and row[idx[name]] not in (None, ''):
                return row[idx[name]]
        return None
    groups = {}
    dates = []
    invalid_dates = 0
    for row in rows:
        name = text(value(row, '건축물명칭', '건물명'))
        other = text(value(row, '기타용도'))
        dong_name = text(value(row, '동명칭'))
        target = ' '.join([name, other, dong_name])
        explicit = bool(re.search(r'데이터\s*센터|data\s*center|\bA?IDC\b', other, re.I))
        if not curated and not re.search(r'데이터\s*센터|data\s*center|\bA?IDC\b', target, re.I):
            continue
        if not name:
            name = f"{text(value(row, '시군구'))} {text(value(row, '법정동'))} 데이터센터(명칭 미기재)"
        bun, ji, land = value(row, '번'), value(row, '지'), value(row, '대지구분')
        # User's appended permit rows have land type in the original lot column.
        shifted = text(bun) in ('대지', '산', '블록') and '대지구분' not in idx
        if shifted:
            land, bun, ji = bun, row[idx['번'] + 1], row[idx['지'] + 1]
        kind = construction_type(None if shifted else value(row, '건축구분'), value(row, '허가구분') or (row[7] if shifted else ''))
        location = row_location(value(row, '시도'), value(row, '시군구'), value(row, '법정동'), bun, ji, text(land))
        permit = date_info(value(row, '허가일'))
        start = date_info(value(row, '착공처리일', '착공일(착공Raw_확정)', '착공일'))
        complete = date_info(value(row, '사용승인일'))
        actual_dates = {'permit': permit, 'start': start, 'completion': complete}
        for d in actual_dates.values():
            if d['text'] and d['kind'] != 'confirmed':
                invalid_dates += 1
        focus = start if starts else permit
        if focus['kind'] == 'confirmed':
            dates.append(focus['date'])
        pk = text(value(row, '건축인허가번호', '인허가번호'))
        fallback = '|'.join([address_key(location) or compact(location), permit['date'] or '', kind])
        key = pk or stable_id(fallback, 'permit')
        item = groups.setdefault(key, {'id': stable_id(key, 'obs'), 'permit_ids': [pk] if pk else [],
            'name': name, 'location': location, 'address_key': address_key(location), 'type': kind,
            'dates': actual_dates, 'area': number(value(row, '연면적(㎡)')), 'area_variants': [],
            'classification': 'confirmed' if explicit or curated else 'candidate',
            'classification_reason': '기존 Raw 분류' if curated else '기타용도에 데이터센터 명시' if explicit else '건물·동 명칭 키워드',
            'sources': [source], 'building_ids': [], 'building_names': []})
        for field, d in actual_dates.items():
            if d['kind'] == 'confirmed':
                old = item['dates'][field]
                if old['kind'] != 'confirmed' or d['date'] < old['date']:
                    item['dates'][field] = d
        area = number(value(row, '연면적(㎡)'))
        if area is not None and area not in item['area_variants']:
            item['area_variants'].append(area)
        if len(item['area_variants']) > 1:
            item['area'] = None  # conflicting total areas are never arbitrarily summed/selected
        dong_id = text(value(row, '동별_개요_PK'))
        if dong_id and dong_id not in item['building_ids']:
            item['building_ids'].append(dong_id)
        if dong_name and dong_name not in item['building_names']:
            item['building_names'].append(dong_name)
    return list(groups.values()), {'rows': len(rows), 'centres': len(groups),
        'date_min': min(dates) if dates else None, 'date_max': max(dates) if dates else None,
        'invalid_dates': invalid_dates}


def import_workbook(path, snapshot_date=None):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = wb['데이터센터(정리)']
    facilities = []
    for row in ws.iter_rows(min_row=2):
        vals = [c.value for c in row]
        if not vals[2] or not vals[1]:
            continue
        raw = dict(zip(['type','location','name','permit','start','alteration','completion','area','capacity','it_load','owner','tenant','contractor','notes','url'], vals))
        ident = '|'.join(text(raw[k]) for k in ['location','name','permit','alteration'])
        item = {'id': stable_id(ident), 'origin': 'baseline', 'source_row': row[0].row,
                'name': text(raw['name']), 'location': text(raw['location']), 'type': text(raw['type']) or '미분류',
                'address_key': address_key(raw['location']),
                'dates': {k: date_info(raw[k]) for k in ['permit','start','alteration','completion']},
                'area': number(raw['area']), 'capacity': number(raw['capacity']),
                'capacity_text': text(raw['capacity']), 'it_load': text(raw['it_load']),
                'owner': text(raw['owner']), 'tenant': text(raw['tenant']), 'contractor': text(raw['contractor']),
                'notes': text(raw['notes']), 'url': text(raw['url']),
                'cell_notes': {c.coordinate: c.comment.text for c in row[:15] if c.comment},
                'classification': 'confirmed', 'classification_reason': '기준자료 정리 목록', 'sources': []}
        facilities.append(item)
    raw_observations = []
    for tab, starts in [('데이터센터 허가(Raw)', False), ('데이터센터 착공(Raw)', True)]:
        values = list(wb[tab].values)
        observations, _ = parse_public_rows(values[0], values[1:], {'id': 'baseline-' + ('starts' if starts else 'permits'),
            'title': tab, 'published': snapshot_date or collection_date(), 'kind': 'starts' if starts else 'permits', 'url': ''}, starts=starts, curated=True)
        raw_observations.extend(observations)
    return {'label': '지엔씨_수전용량 · 데이터센터(정리)', 'snapshot_date': snapshot_date or collection_date(),
            'imported': collection_date(), 'facilities': facilities, 'raw_observations': raw_observations,
            'note': '센터·공사 단계별 정리 목록. 수주·금액 탭과 원본 파일은 포함하지 않음.'}


def same_event(a, b):
    if set(a.get('permit_ids', [])) & set(b.get('permit_ids', [])):
        return True
    # Distinct permit IDs in the same PK generation are distinct permits, even
    # when multiple applications share a site/date. Only bridge legacy/HUB eras.
    left, right = a.get('permit_ids', []), b.get('permit_ids', [])
    era = lambda pk: 'hub' if len(pk.rsplit('-', 1)[-1]) >= 20 else 'legacy'
    if left and right and {era(pk) for pk in left} == {era(pk) for pk in right}:
        return False
    return (bool(a.get('address_key')) and a['address_key'] == b.get('address_key')
            and a['dates']['permit']['date'] is not None
            and a['dates']['permit']['date'] == b['dates']['permit']['date'] and a['type'] == b['type'])


def merge_observations(existing, incoming):
    records = copy.deepcopy(existing)
    for new in incoming:
        matches = [old for old in records if same_event(old, new)]
        if len(matches) != 1:
            if not matches:
                records.append(copy.deepcopy(new))
            continue
        old = matches[0]
        new_pub = max(s['published'] for s in new['sources'])
        old_pub = max(s['published'] for s in old['sources'])
        for key in ['sources','permit_ids','building_ids','building_names']:
            old[key] = list({json.dumps(v, sort_keys=True, ensure_ascii=False): v for v in old[key] + new[key]}.values())
        for key, d in new['dates'].items():
            if d['kind'] == 'confirmed' and (old['dates'][key]['kind'] != 'confirmed' or new_pub >= old_pub):
                old['dates'][key] = copy.deepcopy(d)
        if new_pub >= old_pub:
            old['area'] = new['area']
            old['area_variants'] = new['area_variants']
        if new['classification'] == 'confirmed':
            old['classification'] = 'confirmed'
            old['classification_reason'] = new['classification_reason']
    return records


def matches_facility(item, obs):
    permit = item['dates']['permit']['date']
    alteration = item['dates']['alteration']['date']
    op = obs['dates']['permit']['date']
    date_match = op is not None and op in [permit, alteration]
    same_address = bool(item['address_key']) and item['address_key'] == obs['address_key']
    same_name = compact(item['name']) == compact(obs['name'])
    if date_match and (same_address or (same_name and not (item['address_key'] and obs['address_key']))):
        return True
    return same_address and same_name and permit is None and alteration is None


def reconcile(baseline, observations):
    records = copy.deepcopy(baseline['facilities'])
    changes = []
    for obs in observations:
        matched = [r for r in records if r['origin'] == 'baseline' and matches_facility(r, obs)]
        if len(matched) == 1:
            row = matched[0]
            row['sources'] = list({s['id']: s for s in row['sources'] + obs['sources']}.values())
            row.setdefault('permit_ids', [])
            row['permit_ids'] = sorted(set(row['permit_ids'] + obs['permit_ids']))
            for field, new in obs['dates'].items():
                if new['kind'] != 'confirmed':
                    continue
                # An alteration permit must not replace the original facility permit.
                if field == 'permit' and row['dates']['alteration']['date'] == new['date']:
                    continue
                old = row['dates'][field]
                if old.get('date') == new.get('date') and old['kind'] == 'confirmed':
                    continue
                apply = old['kind'] != 'confirmed'
                changes.append({'id': stable_id('|'.join([row['id'], field, text(old['text']), new['date']]), 'change'),
                    'facility_id': row['id'], 'name': row['name'], 'field': field,
                    'before': old['text'] or '미기재', 'after': new['date'], 'applied': apply,
                    'action': ('기존 Raw로 보완' if all(s['id'].startswith('baseline-') for s in obs['sources']) else '공식 자료로 보완') if apply else '기존 값 유지 · 검토', 'sources': obs['sources']})
                if apply:
                    row.setdefault('original_dates', {})[field] = copy.deepcopy(old)
                    row['dates'][field] = copy.deepcopy(new)
            if row['area'] is not None and obs['area'] is not None and abs(row['area'] - obs['area']) > 1:
                changes.append({'id': stable_id(row['id'] + '|area|' + str(obs['area']), 'change'),
                    'facility_id': row['id'], 'name': row['name'], 'field': 'area', 'before': row['area'],
                    'after': obs['area'], 'applied': False, 'action': '기존 면적 유지 · 범위 확인', 'sources': obs['sources']})
            if row['area'] is None and obs['area'] is not None:
                row['area'] = obs['area']
        elif len(matched) > 1:
            changes.append({'id': stable_id(obs['id'] + '|ambiguous', 'change'), 'name': obs['name'],
                'field': 'match', 'before': '', 'after': '기준자료의 복수 행과 일치', 'applied': False,
                'action': '단계·중복 검토', 'sources': obs['sources']})
        else:
            # A new official record remains a distinct permit/phase, not a presumed unique site.
            row = copy.deepcopy(obs)
            row['origin'] = 'official'
            if all(s['id'].startswith('baseline-') for s in row['sources']):
                row['classification'] = 'candidate'
                row['classification_reason'] = '기준 정리 목록과 아직 연결되지 않은 Raw 행'
            row.update({'capacity': None, 'capacity_text': '', 'it_load': '', 'owner': '', 'tenant': '',
                        'contractor': '', 'notes': '', 'url': '', 'cell_notes': {}})
            row['dates']['alteration'] = date_info(None)
            records.append(row)
    # Repeat snapshots/events can converge to the same published record ID.
    changes = list({c['id']: c for c in changes}.values())
    return records, changes


def annual_series(records, field, measure='count'):
    series = defaultdict(lambda: defaultdict(float))
    for r in records:
        d = r['dates'][field]
        if r['classification'] != 'confirmed' or d['kind'] != 'confirmed':
            continue
        value = 1 if measure == 'count' else r.get(measure)
        if value is None:
            continue
        series[r['type']][d['date'][:4] + '-12-31'] += value
    return {name: [[d, int(v) if measure == 'count' else round(v, 2)] for d, v in sorted(points.items())]
            for name, points in sorted(series.items(), key=lambda x: (x[0] != '신축', x[0]))}


def source_provenance(source):
    return 'user_raw' if source['id'].startswith('baseline-') else 'public_download'


def enrich_records(records, research):
    """Evidence stays separate from the user's facts and administrative dates."""
    entries = research.get('records', {})
    for row in records:
        row['sources'] = [{**s, 'provenance': source_provenance(s)} for s in row['sources']]
        if row['id'] in entries:
            row['research'] = copy.deepcopy(entries[row['id']])
    confirmed = [r for r in records if r['classification'] == 'confirmed']
    examined = [r for r in confirmed if r.get('research')]
    return {'checked': research.get('checked'), 'examined': len(examined),
            'stakeholders_found': sum(bool(r['research'].get('stakeholders')) for r in examined),
            'status_updates': sum(bool(r['research'].get('update')) for r in examined),
            'unresolved_owner': sum(r['owner'].strip() in ('', '?', '미확인', '-') and not any(s['role'] == 'owner' for s in r.get('research', {}).get('stakeholders', [])) for r in confirmed)}


def build_documents(baseline, observations, manifest, status=None, research=None):
    records, changes = reconcile(baseline, merge_observations(baseline['raw_observations'], observations))
    if research is None:
        research = json.loads(RESEARCH.read_text(encoding='utf8')) if RESEARCH.exists() else {}
    research_summary = enrich_records(records, research)
    for change in changes:
        change['sources'] = [{**s, 'provenance': source_provenance(s)} for s in change['sources']]
    today = collection_date()
    common = {'source': '기준 정리자료 + 국토교통부 공개 건축통계', 'source_url': PUBLIC_URL,
        'fetched': status.get('checked', today) if status and status.get('ok') else None,
        'updated': max((s['published'] for s in manifest), default=baseline['snapshot_date']),
        'frequency': 'yearly', 'stale_days': 10,
        'note': '추적 목록의 허가·공사 단계 기준. 전국 전체 통계가 아니며, 이름으로만 추출한 후보·예정일·미확인 날짜는 집계에서 제외. 자료가 없는 연도는 0으로 채우지 않음. 당해 연도는 현재까지 확인된 자료만 포함.'}
    if status:
        common['collection_status'] = status
    docs = {}
    for ident, name, field, measure, unit in [
        ('dc_permits', '연도별 확정 허가 · 추적 목록', 'permit', 'count', '건'),
        ('dc_starts', '연도별 확정 착공 · 추적 목록', 'start', 'count', '건'),
        ('dc_started_area', '착공 연도별 연면적 · 확인된 면적', 'start', 'area', '㎡'),
        ('dc_started_capacity', '착공 연도별 수전용량 · 확인된 용량', 'start', 'capacity', 'MW')]:
        pts = annual_series(records, field, measure)
        docs[ident] = {**copy.deepcopy(common), 'id': ident, 'name': name, 'unit': unit,
            'year_labels': True, 'annual_axis': True, 'span_gaps': False, 'change_mode': 'none',
            'series': pts, 'default_series': ['신축'] if '신축' in pts else list(pts)[:1]}
    docs['dc_facilities'] = {**common, 'id': 'dc_facilities', 'name': '국내 데이터센터 · 센터·공사 단계별 현황',
        'records': records, 'changes': changes, 'manifest': manifest, 'research_summary': research_summary,
        'baseline': {k: baseline[k] for k in ['label','snapshot_date','imported']},
        'summary': {'baseline_rows': len(baseline['facilities']), 'tracked_rows': sum(r['classification'] == 'confirmed' for r in records),
                    'candidate_rows': sum(r['classification'] != 'confirmed' for r in records),
                    'applied': sum(c['applied'] for c in changes), 'review': sum(not c['applied'] for c in changes)},
        'series': {}}
    return docs


def write_documents(docs):
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for name, doc in docs.items():
        path = OUTPUT / (name + '.json')
        if doc.get('fetched') is None and path.exists():
            previous = json.loads(path.read_text(encoding='utf8'))
            doc['fetched'] = previous.get('fetched')
        path.write_text(json.dumps(doc, ensure_ascii=False, indent=1) + '\n', encoding='utf8')


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--import-xlsx', type=Path)
    parser.add_argument('--snapshot-date', help='Actual source file modification/snapshot date, YYYY-MM-DD')
    args = parser.parse_args()
    if args.import_xlsx:
        if args.snapshot_date:
            date.fromisoformat(args.snapshot_date)
        baseline = import_workbook(args.import_xlsx, args.snapshot_date)
        BASELINE.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + '\n', encoding='utf8')
        print('Imported curated rows:', len(baseline['facilities']))

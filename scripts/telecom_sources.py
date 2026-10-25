"""Public issuer sources; values in original-currency millions, with exact periods."""
import calendar
import csv
import datetime as dt
import io
import math
import re
from urllib.parse import urljoin

import openpyxl
from bs4 import BeautifulSoup

from fetch_cpu_companies import parse_facts, parse_annual


def reporting_period(year, quarter=None, fiscal_start=1):
    start_year=year-(fiscal_start!=1)
    start=dt.date(start_year,fiscal_start,1)
    month_index=start.year*12+start.month-1+(quarter-1)*3 if quarter else start.year*12+start.month-1
    begin=dt.date(month_index//12,month_index%12+1,1)
    next_index=month_index+(3 if quarter else 12)
    end=dt.date(next_index//12,next_index%12+1,1)-dt.timedelta(days=1)
    return dict(kind='quarter' if quarter else 'annual',year=year,quarter=quarter,start=begin.isoformat(),date=end.isoformat(),period=f'FY{year}'+(f' Q{quarter}' if quarter else ''),refs={})


def sec_accounts(payload, meta, asof):
    rows=[]
    for kind,items in [('quarter',parse_facts(payload,meta,asof)),('annual',parse_annual(payload,meta,asof))]:
        for r in items:
            if r['date']<'2020-01-01':continue
            if not r.get('start'):
                r=dict(r,start=next(iter(r['refs'].values()))['start'])
            rows.append(dict(r,kind=kind,period=f"FY{r['year']}"+(f" Q{r['quarter']}" if kind=='quarter' else '')))
    return sorted(rows,key=lambda r:(r['date'],r['kind']))


def number(value):
    if isinstance(value,bool):return None
    if isinstance(value,(int,float)):return float(value) if math.isfinite(value) else None
    text=str(value or '').strip().replace(',','').replace('−','-')
    if text.startswith('(') and text.endswith(')'):text='-'+text[1:-1]
    try:return float(text) if text and math.isfinite(float(text)) else None
    except ValueError:return None


def dart_report(items,year,quarter,corp):
    if any(r.get('corp_code',corp)!=corp for r in items):raise ValueError('DART entity mismatch')
    fields={}
    for key,tags,names,divs in [
        ('revenue',('ifrs-full_Revenue','ifrs-full_RevenueFromContractsWithCustomers'),('매출액','영업수익','수익(매출액)','매출'),('IS','CIS')),
        ('profit',('dart_OperatingIncomeLoss','ifrs-full_ProfitLossFromOperatingActivities'),('영업이익','영업이익(손실)','영업손익'),('IS','CIS')),
        ('capex',('ifrs-full_PurchaseOfPropertyPlantAndEquipment','ifrs-full_PaymentsToAcquirePropertyPlantAndEquipment'),('유형자산의취득','유형자산취득','유형자산의취득으로인한현금유출','유형자산의취득으로인한현금유출액'),('CF',))]:
        candidates=[r for r in items if r.get('sj_div') in divs and (r.get('account_id') in tags or r.get('account_nm','').replace(' ','') in names)]
        if not candidates:continue
        item=min(candidates,key=lambda r:0 if r.get('account_id') in tags else 1)
        value=number(item.get('thstrm_amount'))
        ytd=value if key=='capex' or quarter in (1,4) else number(item.get('thstrm_add_amount'))
        if key=='capex':value=abs(value) if value is not None else None;ytd=value
        acc=item['rcept_no']
        fields[key]=dict(value=value/1e6 if value is not None else None,ytd=ytd/1e6 if ytd is not None else None,
                         url=f'https://dart.fss.or.kr/dsaf001/main.do?rcpNo={acc}',filed=f'{acc[:4]}-{acc[4:6]}-{acc[6:8]}',tag=item['account_id'])
    if not all(k in fields for k in ('revenue','profit')):raise ValueError('missing consolidated income statement')
    return dict(year=year,quarter=quarter,fields=fields)


def dart_accounts(reports,asof):
    index={(r['year'],r['quarter']):r for r in reports};rows=[]
    for (year,q),report in sorted(index.items()):
        if q==4:
            annual=reporting_period(year)
            for key,f in report['fields'].items():
                if f['ytd'] is not None and f['filed']<=asof:
                    annual[key]=f['ytd'];annual['refs'][key]=dict(f,calculation='reported annual')
            if annual['refs']:rows.append(annual)
        row=reporting_period(year,q)
        for key,f in report['fields'].items():
            if f['filed']>asof:continue
            if q!=4 and key!='capex' and f['value'] is not None:
                row[key]=f['value'];row['refs'][key]=dict(f,calculation='reported quarter');continue
            previous=index.get((year,q-1),{}).get('fields',{}).get(key)
            if q==1 and f['ytd'] is not None:
                row[key]=f['ytd'];row['refs'][key]=dict(f,calculation='reported quarter')
            elif previous and previous.get('ytd') is not None and f['ytd'] is not None and previous['tag']==f['tag'] and previous['filed']<=f['filed']:
                value=f['ytd']-previous['ytd']
                if key!='profit' and value<0:continue
                row[key]=value;row['refs'][key]=dict(f,calculation='same fiscal-year YTD minus prior YTD',components=[f,previous])
        if row['date']<=asof and any(k in row for k in ('revenue','profit','capex')):rows.append(row)
    return sorted(rows,key=lambda r:(r['date'],r['kind']))


def workbook(raw):
    return openpyxl.load_workbook(io.BytesIO(raw),data_only=True)


def columns(sheet, header, year_row):
    """Only actual Q1..Q4/full-year columns. Forecast columns never form a group."""
    starts=[c.column for c in sheet[header] if re.search(r'(?:^|\W)(?:1Q|Q1)(?:\W|$)',str(c.value or '').strip())]
    result=[]
    for col in starts:
        candidates=[c.value for row in sheet.iter_rows(min_row=year_row,max_row=year_row,min_col=col,max_col=min(col+4,sheet.max_column)) for c in row if re.fullmatch(r'FY\s?\d{2,4}',str(c.value or '').strip())]
        if len(candidates)!=1:continue
        fiscal=int(re.search(r'\d+',str(candidates[0]))[0]);fiscal=fiscal+2000 if fiscal<100 else fiscal
        for offset,q in [(0,1),(1,2),(2,3),(3,4),(4,None)]:
            if col+offset>sheet.max_column:continue
            label=str(sheet.cell(header,col+offset).value or '')
            if 'forecast' in label.lower() or '予想' in label:continue
            if offset==4 and not any(t in label.lower() for t in ['full year','results','通期']):continue
            result.append((col+offset,fiscal+1,q))  # issuer FY is starting-year basis
    return result


def merge_value(rows,period,key,value,url,location,asof,scale=1000,absolute=False):
    if value is None or period['date']>asof or period['date']<'2020-01-01':return
    k=period['kind'],period['date'];row=rows.setdefault(k,dict(period,refs={}))
    row[key]=(abs(value) if absolute else value)*scale
    row['refs'][key]=dict(url=url,location=location,calculation='reported quarter' if period['kind']=='quarter' else 'reported annual')


def parse_ntt(raw,url,asof):
    w=workbook(raw);rows={}
    def matches(value,term):
        return any(line.split('/')[-1].strip().casefold()==term.casefold() for line in str(value).splitlines())
    income=next((s for s in w if any(matches(c.value,'Operating revenues') for row in list(s)[:35] for c in row[:8]) and any(matches(c.value,'Operating profit') for row in list(s)[:35] for c in row[:8])),None)
    if income is None:raise ValueError('NTT income sheet changed')
    for key,term in [('revenue','Operating revenues'),('profit','Operating profit')]:
        candidates=[r for r in income if any(matches(c.value,term) for c in r[:8])]
        if not candidates:raise ValueError('NTT metric layout changed')
        r=candidates[0]
        for col,year,q in columns(income,5,4):merge_value(rows,reporting_period(year,q,4),key,number(r[col-1].value),url,f'1!{r[col-1].coordinate}',asof)
    cap_sheets=[s for s in w if 'capital investment' in ' '.join(str(c.value) for c in s[1] if c.value is not None).lower()]
    if len(cap_sheets)!=1:raise ValueError('NTT CAPEX sheet changed')
    s=cap_sheets[0]
    cap=[r for r in s if any(matches(c.value,'Capital Investment') for c in r[:8])]
    if not cap:raise ValueError('NTT CAPEX row changed')
    for col,year,q in columns(s,5,4):merge_value(rows,reporting_period(year,q,4),'capex',number(cap[0][col-1].value),url,f'{s.title}!{cap[0][col-1].coordinate}',asof)
    return list(rows.values())


def parse_softbank(raw,url,asof):
    w=workbook(raw);rows={}
    income=w['Consolidated_R&OP_1'];cash=w['Consolidated_R&OP_3']
    for s,key,term,absolute in [(income,'revenue','売上高\nRevenue',False),(income,'profit','営業利益\nOperating income',False),(cash,'capex','設備支出 (C)\nCapital expenditures',True)]:
        found=[r for r in s if any(str(c.value).strip()==term for c in r[:5])]
        if len(found)!=1:raise ValueError(f'SoftBank {key} row changed')
        r=found[0]
        for col,year,q in columns(s,6,5):merge_value(rows,reporting_period(year,q,4),key,number(r[col-1].value),url,f'{s.title}!{r[col-1].coordinate}',asof,absolute=absolute)
    return list(rows.values())


def parse_kddi(raw,url,asof):
    w=workbook(raw);rows={};s=w.worksheets[0]
    if 'Financial Highlights' not in s.title:raise ValueError('KDDI sheet changed')
    for col in range(18,s.max_column+1):  # IFRS block; avoid duplicate old J-GAAP years
        match=re.fullmatch(r'FY(20\d{2})\.3',str(s.cell(5,col).value or ''))
        if not match:continue
        for key,row,term in [('revenue',6,'Operating revenue'),('profit',9,'Operating income'),('capex',16,'Capex (cash flow basis)')]:
            if term not in str(s.cell(row,1).value):raise ValueError('KDDI annual layout changed')
            merge_value(rows,reporting_period(int(match[1]),None,4),key,number(s.cell(row,col).value),url,f'{s.title}!{s.cell(row,col).coordinate}',asof,scale=1)
    s=w['Quarterly Basis (PL)']
    for r in s:
        if not any(re.fullmatch(r'FY20\d{2}\.3',str(c.value or '')) for c in r):continue
        for c in r:
            m=re.fullmatch(r'FY(20\d{2})\.3',str(c.value or ''))
            if not m:continue
            for key,offset,term in [('revenue',2,'Operating revenue'),('profit',3,'Operating income')]:
                metric_row=r[0].row+offset
                # Since FY2027.3 this row is ADJUSTED operating income. Do not
                # silently join it to the historical IFRS operating profit.
                label=str(s.cell(metric_row,1).value or '')
                if term not in label or 'Adjusted' in label:continue
                for q in range(1,5):
                    cell=s.cell(metric_row,c.column+q-1)
                    merge_value(rows,reporting_period(int(m[1]),q,4),key,number(cell.value),url,f'{s.title}!{cell.coordinate}',asof)
    return list(rows.values())


def softbank_link(raw,asof='9999-12-31'):
    rows=list(csv.reader(io.StringIO(raw.decode('utf-8-sig')),delimiter='\t'))
    for row in rows:
        if not row or not re.fullmatch(r'20\d{2}',row[0]):continue
        if len(row)<3 or not re.fullmatch(r'\d{8}',row[2]) or row[2]>asof.replace('-',''):continue
        for cell in row:
            if '.xlsx' in cell:return urljoin('https://www.softbank.jp',cell)
    raise ValueError('SoftBank published workbook missing')

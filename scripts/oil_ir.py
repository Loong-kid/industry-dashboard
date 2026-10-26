"""Quarterly oil financial statements from original issuer PDFs and workbooks."""
import datetime as dt
import json
from pathlib import Path
import re
import openpyxl
import pymupdf
from oil_sources import pdf_row,pdf_lines,quarter_end,NUMBER,number

def ref(url,start,end,location,filed):
    page=re.search(r'PDF p\.(\d+)',location)
    if page:url+='#page='+page[1]
    return dict(url=url,start=start,end=end,filed=filed,tag=location,components=[dict(url=url,start=start,end=end,filed=filed,tag=location)],calculation='original issuer financial statement')

def tte_pdf(path,url,filed):
    y,q=map(int,re.search(r'tte-(\d+)-q(\d)',path.name).groups());end=quarter_end(y,q);start=f'{y}-{(q-1)*3+1:02d}-01';result=[]
    d=pymupdf.open(path)
    for index,page in enumerate(d):
        text=page.get_text();upper=text.upper()
        if 'CONSOLIDATED STATEMENT OF INCOME' not in upper and 'CONSOLIDATED STATEMENT OF CASH FLOW' not in upper:continue
        # Quarter and full year statements have separate page headings.
        annual=q==4 and bool(re.search(r'(?:Year|year)\s+(?:\n)?\s*'+str(y),page.get_text(sort=True)))
        quarterly=index==0 or ('CASH FLOW' in upper and re.search(r'(?:1st|2nd|3rd|4th)\s+quarter',page.get_text(sort=True),re.I))
        if not (quarterly or annual):continue
        period_start=f'{y}-01-01' if annual else start
        location=f'PDF p.{index+1} · '+('annual' if annual else 'quarter')
        values={}
        if 'STATEMENT OF INCOME' in upper:
            patterns={'revenue':r'^Revenues? from sales\b','purchase':r'^Purchases, net of inventory variation\b','opex':r'^Other operating expenses\b','explore':r'^Exploration costs\b','dda':r'^Depreciation, depletion and (?:impairment|amortization).*mineral interests','other_income':r'^Other income\b','other_expense':r'^Other expenses?\b'}
            for field,pattern in patterns.items():
                v=pdf_row(page,pattern)
                if v:values[field]=v
            if 'revenue' not in values:continue
            periods=[(y-i,4) for i in range(len(values['revenue']))] if annual else [(y,q),(y-(q==1),4 if q==1 else q-1),(y-1,q)]
            for column,(year,quarter) in enumerate(periods):
                if year<2014 or column>=len(values['revenue']):continue
                actual_end=quarter_end(year,quarter);actual_start=f'{year}-01-01' if annual else f'{year}-{(quarter-1)*3+1:02d}-01'
                row=dict(date=actual_end,year=year,refs={},**({} if annual else dict(quarter=quarter,start=actual_start)))
                row['revenue']=values['revenue'][column];row['refs']['revenue']=ref(url,actual_start,actual_end,location+' · Revenues from sales (excise excluded)',filed)
                if len(values)==len(patterns) and all(column<len(v) for v in values.values()):
                    signed={field:v[column] for field,v in values.items()}
                    row['profit']=sum(signed.values());row['refs']['profit']=ref(url,actual_start,actual_end,location+' · Sales − purchases − operating expenses − exploration − depreciation + other income − other expense',filed)
                    row['refs']['profit']['calculation']='calculated operating subtotal from the signed statement rows';row['refs']['profit']['inputs']=signed
                result.append((annual,row))
        else:
            v=pdf_row(page,r'^Intangible assets and property, plant and equipment additions\b')
            if v:
                periods=[(y-i,4) for i in range(len(v))] if annual else [(y,q),(y-(q==1),4 if q==1 else q-1),(y-1,q)]
                for column,(year,quarter) in enumerate(periods):
                    if year<2014 or column>=len(v):continue
                    actual_end=quarter_end(year,quarter);actual_start=f'{year}-01-01' if annual else f'{year}-{(quarter-1)*3+1:02d}-01'
                    row=dict(date=actual_end,year=year,capex=abs(v[column]),refs={'capex':ref(url,actual_start,actual_end,location+' · Intangible assets and PP&E additions',filed)},**({} if annual else dict(quarter=quarter,start=actual_start)));result.append((annual,row))
    return result

def shell_2020_pdf(path,url):
    d=pymupdf.open(path);income=d[1];cash=d[3];fields={}
    for field,page,pattern in [('revenue',income,r'^Revenue\b'),('pbt',income,r'^Income/(?:\(loss\))? before taxation\b'),('interest',income,r'^Interest expense\b'),('other',income,r'^Interest and other income\b'),('capex',cash,r'^Capital expenditure\b')]:
        values=pdf_row(page,pattern)
        if values is not None:fields[field]=values
    if len(fields)!=5 or any(len(values)!=12 for values in fields.values()):raise ValueError('Shell 2020 databook columns changed')
    periods=[(True,2017,4),(True,2018,4)]+[(q==5,y,min(q,4)) for y in [2019,2020] for q in range(1,6)]
    result=[]
    for i,(annual,y,q) in enumerate(periods):
        end=quarter_end(y,q);start=f'{y}-01-01' if annual else f'{y}-{(q-1)*3+1:02d}-01';v={k:vs[i] for k,vs in fields.items()}
        row=dict(date=end,year=y,revenue=v['revenue'],profit=v['pbt']+v['interest']-v['other'],capex=abs(v['capex']),refs={},**({} if annual else dict(quarter=q,start=start)))
        for field in ['revenue','profit','capex']:row['refs'][field]=ref(url,start,end,'PDF p.'+('4' if field=='capex' else '2'),'2021-02-04')
        row['refs']['profit'].update(calculation='PBT + interest expense − interest and other income',inputs=v)
        result.append((annual,row))
    return result

def equinor_pdf(path,url,filed):
    y,q=map(int,re.search(r'equinor-(\d+)-q(\d)',path.name).groups())
    if y<2016:return [] # USD comparative workbook supplies 2014–15; no NOK splice.
    end=quarter_end(y,q);start=f'{y}-{(q-1)*3+1:02d}-01';result=[]
    d=pymupdf.open(path)
    income_seen=False;cash_seen=False
    for index,page in enumerate(d):
        text=page.get_text();upper=text.upper()
        income='CONSOLIDATED STATEMENT OF INCOME' in upper and not income_seen
        cash='CONSOLIDATED STATEMENT OF CASH FLOWS' in upper and not cash_seen
        if not (income or cash):continue
        patterns={'revenue':r'\bRevenues\b','profit':r'\bNet operating income(?:/\(loss\))?\b'} if income else {'capex':r'\bCapital expenditures and investments'}
        xmax=page.rect.width/2 if page.rect.width>page.rect.height else None
        columns=[];period_columns=[]
        for line in pdf_lines(page):
            if xmax:line=[w for w in line if w[2]<xmax]
            dates=[];years=[]
            for i,w in enumerate(line):
                if re.fullmatch(r'Q[1-4]',w[4]) and i+1<len(line) and re.fullmatch(r'20\d{2}\*?',line[i+1][4]):
                    dates.append((int(w[4][1]),int(line[i+1][4][:4]),(w[0]+line[i+1][2])/2))
                if re.fullmatch(r'20\d{2}\*?',w[4]) and (i==0 or not re.fullmatch(r'Q[1-4]',line[i-1][4])):years.append((int(w[4][:4]),(w[0]+w[2])/2))
            if len(dates)>=3 and dates[0][:2]==(q,y):
                columns=[c[2] for c in dates]+[c[1] for c in years]
                period_columns=[(False,year,quarter,center) for quarter,year,center in dates]
                if q==4:period_columns += [(True,year,4,center) for year,center in years]
                break
        if not columns:continue
        extracted={}
        for field,pattern in patterns.items():
            values={}
            for line in pdf_lines(page):
                if xmax:line=[w for w in line if w[2]<xmax]
                if not re.search(pattern,' '.join(w[4] for w in line),re.I):continue
                for annual,year,quarter,center in period_columns:
                    name=(annual,year,quarter)
                    if center is None:continue
                    ordered=sorted(columns);position=ordered.index(center)
                    low=(ordered[position-1]+center)/2 if position else center-(ordered[1]-center)/2
                    high=(ordered[position+1]+center)/2 if position+1<len(ordered) else center+(center-ordered[position-1])/2
                    tokens=NUMBER.findall(' '.join(w[4] for w in line if low<=(w[0]+w[2])/2<high))
                    if len(tokens)==1:values[name]=number(tokens[0])
                if values:break
            if (False,y,q) in values:extracted[field]=values
        if len(extracted)!=len(patterns):continue
        if income:income_seen=True
        else:cash_seen=True
        for annual,year,quarter,_ in period_columns:
            if year<2014:continue
            position=(annual,year,quarter)
            actual_end=quarter_end(year,quarter);actual_start=f'{year}-01-01' if annual else f'{year}-{(quarter-1)*3+1:02d}-01'
            row=dict(date=actual_end,year=year,refs={},**({} if annual else dict(quarter=quarter,start=actual_start)))
            for field,values in extracted.items():
                if position not in values:continue
                value=abs(values[position]) if field=='capex' else values[position]
                row[field]=value;row['refs'][field]=ref(url,actual_start,actual_end,f'PDF p.{index+1} · '+patterns[field],filed)
            result.append((annual,row))
        if income_seen and cash_seen:break
    return result

def shell_workbook(path,url,filed):
    w=openpyxl.load_workbook(path,data_only=True,read_only=True);result=[]
    income=w['Statement of Income'];cash=w['Cash Flows']
    def row(sheet,label):
        for r in sheet.iter_rows(max_col=36):
            if any(isinstance(c.value,str) and re.match(label,c.value.strip(),re.I) for c in r[:3]):return r
        raise ValueError('Shell workbook row missing: '+label)
    rows={field:row(sheet,pattern) for field,sheet,pattern in [('revenue',income,r'^Revenue$'),('pbt',income,r'^Income/\(loss\) before taxation$'),('interest',income,r'^Interest expense$'),('other',income,r'^Interest and other income'),('capex',cash,r'^Capital expenditure$')]}
    for cell in income[7]:
        match=re.fullmatch(r'Q([1-4]) (20\d{2})',str(cell.value));annual=bool(re.fullmatch(r'20\d{2}',str(cell.value)))
        if not match and not annual:continue
        y=int(match[2]) if match else int(cell.value);q=int(match[1]) if match else 4
        # Rightmost "2026" column is H1, not a full-year observation.
        if annual and y>=int(filed[:4]):continue
        v={field:cells[cell.column-1].value for field,cells in rows.items()}
        if not all(isinstance(value,(int,float)) for value in v.values()):continue
        end=quarter_end(y,q);start=f'{y}-01-01' if annual else f'{y}-{(q-1)*3+1:02d}-01'
        r=dict(date=end,year=y,revenue=v['revenue'],profit=v['pbt']+v['interest']-v['other'],capex=abs(v['capex']),refs={},**({} if annual else dict(quarter=q,start=start)))
        for field in ['revenue','profit','capex']:r['refs'][field]=ref(url,start,end,('Statement of Income' if field!='capex' else 'Cash Flows')+' · '+cell.coordinate,filed)
        r['refs']['profit']['calculation']='Income before taxation + interest expense − interest and other income';r['refs']['profit']['inputs']=v
        result.append((annual,r))
    return result

def equinor_usd_history(path,url):
    """Company's March 2016 USD restatement, not a dashboard FX conversion."""
    w=openpyxl.load_workbook(path,data_only=True);result=[]
    for title,fields,header in [('Cons_Income_QTD',{'revenue':'Revenues','profit':'Net operating income'},3),('Cons_CashFlow_QTD',{'capex':'Capital expenditures and investments'},2),('Cons_Income_YTD',{'revenue':'Revenues','profit':'Net operating income'},3),('Cons_CashFlow_YTD',{'capex':'Capital expenditures and investments'},2)]:
        s=w[title];annual='YTD' in title
        for cell in s[header]:
            m=re.search(r'Q([1-4]) (201[45])',str(cell.value))
            if not m:continue
            q,y=int(m[1]),int(m[2])
            if annual and q!=4:continue
            r=dict(date=quarter_end(y,q),year=y,refs={},**({} if annual else dict(quarter=q,start=f'{y}-{(q-1)*3+1:02d}-01')))
            for field,label in fields.items():
                for row in s:
                    if str(row[0].value).strip()!=label:continue
                    v=row[cell.column-1].value
                    if not isinstance(v,(float,int)):continue
                    r[field]=round(abs(v) if field=='capex' else v,3)
                    r['refs'][field]=ref(url,f'{y}-01-01' if annual else r['start'],r['date'],title+'!'+row[cell.column-1].coordinate,'2016-03-31')
            result.append((annual,r))
    return result
